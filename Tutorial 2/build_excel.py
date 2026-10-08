"""Builds Tutorial2_Q2_Q3_Solutions.xlsx - live, Solver-ready sheets for Q2 and Q3."""
import json
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.chart.marker import Marker
from openpyxl.utils import get_column_letter

from tut2_data import CARBON, SILICA, CH4, CO2

RES = json.load(open("tut2_results.json"))
TS = [298, 308, 323]

H1 = Font(bold=True, size=14, color="FFFFFF")
HDR = Font(bold=True, size=10)
BOLD = Font(bold=True)
NOTE = Font(italic=True, size=9, color="555555")
FILL_T = PatternFill("solid", fgColor="1F4E79")
FILL_P = PatternFill("solid", fgColor="FFF2CC")   # yellow = Solver decision variables
FILL_O = PatternFill("solid", fgColor="D9EAD3")   # green  = Solver objective
FILL_H = PatternFill("solid", fgColor="DDEBF7")
THIN = Border(*[Side(style="thin", color="BFBFBF")] * 4)


def title(ws, text, span=6):
    ws["A1"] = text
    ws["A1"].font = H1
    for c in range(1, span + 1):
        ws.cell(1, c).fill = FILL_T
    ws.row_dimensions[1].height = 22


def note(ws, row, text, span=8):
    ws.cell(row, 1, text).font = NOTE
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def param_block(ws, row, names, values, fmt="0.000000"):
    """Yellow decision-variable cells. Returns {name: 'B$n'} absolute refs."""
    ws.cell(row, 1, "Parameters (Solver: By Changing Variable Cells)").font = BOLD
    refs = {}
    for i, (n, v) in enumerate(zip(names, values)):
        r = row + 1 + i
        ws.cell(r, 1, n).font = HDR
        c = ws.cell(r, 2, round(float(v), 8))
        c.fill, c.border, c.number_format = FILL_P, THIN, fmt
        refs[n] = f"$B${r}"
    return refs, row + 1 + len(names)


def scatter(ws, anchor, title_txt, x_ref, y_exp, y_cal, xt, yt, logx=False):
    ch = ScatterChart()
    ch.title, ch.x_axis.title, ch.y_axis.title = title_txt, xt, yt
    ch.style, ch.height, ch.width = 13, 8.5, 14
    ch.x_axis.majorGridlines = None
    if logx:
        ch.x_axis.scaling.logBase = 10
    s1 = Series(y_exp, x_ref, title="experimental")
    s1.marker = Marker(symbol="circle", size=6)
    s1.graphicalProperties.line.noFill = True
    s2 = Series(y_cal, x_ref, title="model")
    s2.marker = Marker(symbol="none")
    s2.graphicalProperties.line.width = 20000
    ch.series.append(s1)
    ch.series.append(s2)
    ws.add_chart(ch, anchor)


# =====================================================================
# Q2 - single-site Langmuir
# =====================================================================
def sheet_q2(wb, name, data, key, pun="mmHg", qun="mmol/g"):
    ws = wb.create_sheet(name)
    widths(ws, {"A": 30, "B": 14, "C": 14, "D": 14, "E": 14, "F": 14})
    r = RES["Q2"][key]
    title(ws, f"Q2 - Langmuir isotherm, CO2 on {name.split('_')[-1]} at 25 °C")
    note(ws, 2, "Model:  q = qmax*K*p/(1+K*p).   Yellow = Solver variables, green = objective "
                "(minimise SSE).  Solver: GRG Nonlinear, Min, variables qmax & K, both >= 0.")

    refs, last = param_block(ws, 4, ["qmax  (%s)" % qun, "K  (1/%s)" % pun],
                             [r["qmax"], r["K"]])
    ws.cell(5, 3, "<- maximum adsorbate loading").font = NOTE
    ws.cell(6, 3, "<- adsorption equilibrium constant").font = NOTE

    hr = last + 2
    hdrs = [f"p ({pun})", f"q_exp ({qun})", f"q_calc ({qun})", "residual", "residual^2", "theta"]
    for j, h in enumerate(hdrs, 1):
        c = ws.cell(hr, j, h)
        c.font, c.fill, c.border = HDR, FILL_H, THIN
    for i, (p, q) in enumerate(data):
        rw = hr + 1 + i
        ws.cell(rw, 1, p).number_format = "0.0"
        ws.cell(rw, 2, q).number_format = "0.0000"
        ws.cell(rw, 3, f"={refs['qmax  (%s)' % qun]}*{refs['K  (1/%s)' % pun]}*A{rw}"
                       f"/(1+{refs['K  (1/%s)' % pun]}*A{rw})").number_format = "0.0000"
        ws.cell(rw, 4, f"=C{rw}-B{rw}").number_format = "0.0000"
        ws.cell(rw, 5, f"=D{rw}^2").number_format = "0.00E+00"
        ws.cell(rw, 6, f"=C{rw}/{refs['qmax  (%s)' % qun]}").number_format = "0.0000"
        for j in range(1, 7):
            ws.cell(rw, j).border = THIN
    n = len(data)
    fr, lr = hr + 1, hr + n

    sr = lr + 2
    ws.cell(sr, 1, "SSE  (Solver: Set Objective -> Min)").font = BOLD
    o = ws.cell(sr, 2, f"=SUM(E{fr}:E{lr})")
    o.fill, o.border, o.number_format = FILL_O, THIN, "0.000000"
    ws.cell(sr + 1, 1, "SST").font = HDR
    ws.cell(sr + 1, 2, f"=DEVSQ(B{fr}:B{lr})").number_format = "0.000000"
    ws.cell(sr + 2, 1, "R^2  = 1 - SSE/SST").font = HDR
    ws.cell(sr + 2, 2, f"=1-B{sr}/B{sr+1}").number_format = "0.00000"
    ws.cell(sr + 3, 1, "RMSE").font = HDR
    ws.cell(sr + 3, 2, f"=SQRT(B{sr}/{n})").number_format = "0.00000"
    ws.cell(sr + 4, 1, "AARD (%)").font = HDR
    ws.cell(sr + 4, 2, f"=100*AVERAGE(ABS(D{fr}:D{lr})/B{fr}:B{lr})").number_format = "0.00"
    ws.cell(sr + 4, 3, "array formula - confirm with Ctrl+Shift+Enter in older Excel").font = NOTE

    lr2 = sr + 6
    ws.cell(lr2, 1, "Linearised check:  p/q = 1/(qmax*K) + p/qmax").font = BOLD
    ws.cell(lr2 + 1, 1, "slope = 1/qmax").font = HDR
    ws.cell(lr2 + 1, 2, f"=SLOPE(G{fr}:G{lr},A{fr}:A{lr})").number_format = "0.000000"
    ws.cell(lr2 + 2, 1, "intercept = 1/(qmax*K)").font = HDR
    ws.cell(lr2 + 2, 2, f"=INTERCEPT(G{fr}:G{lr},A{fr}:A{lr})").number_format = "0.000000"
    ws.cell(lr2 + 3, 1, "qmax (linear)").font = HDR
    ws.cell(lr2 + 3, 2, f"=1/B{lr2+1}").number_format = "0.0000"
    ws.cell(lr2 + 4, 1, "K (linear)").font = HDR
    ws.cell(lr2 + 4, 2, f"=B{lr2+1}/B{lr2+2}").number_format = "0.000000"
    ws.cell(hr, 7, "p/q").font = HDR
    ws.cell(hr, 7).fill = FILL_H
    for i in range(n):
        rw = fr + i
        ws.cell(rw, 7, f"=A{rw}/B{rw}").number_format = "0.000"

    x = Reference(ws, min_col=1, min_row=fr, max_row=lr)
    scatter(ws, "I4", f"CO2 on {key} - Langmuir fit", x,
            Reference(ws, min_col=2, min_row=hr, max_row=lr),
            Reference(ws, min_col=3, min_row=hr, max_row=lr),
            f"p ({pun})", f"q ({qun})")
    return ws


# =====================================================================
# Q3 - Toth
# =====================================================================
def sheet_toth(wb, gas, sets):
    ws = wb.create_sheet(f"Q3_{gas}_Toth")
    widths(ws, {"A": 26, **{get_column_letter(c): 12 for c in range(2, 20)}})
    r = RES["Q3"][gas]["toth"]
    title(ws, f"Q3 - Toth isotherm, {gas} on zeolite 13X", 12)
    note(ws, 2, "Model:  q = qmax*K*p / [1+(K*p)^n]^(1/n).   qmax and n are shared by all three "
                "isotherms (temperature-independent); one K per temperature.  Solver: GRG "
                "Nonlinear, Min total SSE, variables qmax, n, K298, K308, K323 (all >= 0).", 12)

    names = ["qmax  (mol/kg)", "n  (-)"] + [f"K{T}  (1/MPa)" for T in TS]
    vals = [r["qmax"], r["n"]] + [r["K"][str(T)] for T in TS]
    refs, last = param_block(ws, 4, names, vals)

    hr = last + 2
    blocks = []
    for bi, T in enumerate(TS):
        c0 = 1 + bi * 5          # 1, 6, 11
        ws.cell(hr - 1, c0, f"T = {T} K").font = BOLD
        ws.cell(hr - 1, c0).fill = FILL_H
        for j, h in enumerate(["P (MPa)", "q_exp", "q_calc", "resid", "resid^2"]):
            c = ws.cell(hr, c0 + j, h)
            c.font, c.fill, c.border = HDR, FILL_H, THIN
        data = sets[T]
        for i, (p, q) in enumerate(data):
            rw = hr + 1 + i
            L = get_column_letter
            a, b, cc, d = L(c0), L(c0 + 1), L(c0 + 2), L(c0 + 3)
            ws.cell(rw, c0, p).number_format = "0.00000"
            ws.cell(rw, c0 + 1, q).number_format = "0.000"
            ws.cell(rw, c0 + 2,
                    f"={refs['qmax  (mol/kg)']}*{refs[f'K{T}  (1/MPa)']}*{a}{rw}"
                    f"/(1+({refs[f'K{T}  (1/MPa)']}*{a}{rw})^{refs['n  (-)']})"
                    f"^(1/{refs['n  (-)']})").number_format = "0.0000"
            ws.cell(rw, c0 + 3, f"={cc}{rw}-{b}{rw}").number_format = "0.0000"
            ws.cell(rw, c0 + 4, f"={d}{rw}^2").number_format = "0.00E+00"
            for j in range(5):
                ws.cell(rw, c0 + j).border = THIN
        blocks.append((c0, hr + 1, hr + len(data)))

    sr = hr + max(len(sets[T]) for T in TS) + 2
    L = get_column_letter
    for bi, (c0, f_, l_) in enumerate(blocks):
        ws.cell(sr, c0, f"SSE {TS[bi]} K").font = HDR
        ws.cell(sr, c0 + 1, f"=SUM({L(c0+4)}{f_}:{L(c0+4)}{l_})").number_format = "0.00000"
        ws.cell(sr + 1, c0, f"R^2 {TS[bi]} K").font = HDR
        ws.cell(sr + 1, c0 + 1,
                f"=1-{L(c0+1)}{sr}/DEVSQ({L(c0+1)}{f_}:{L(c0+1)}{l_})").number_format = "0.00000"
    tr = sr + 3
    ws.cell(tr, 1, "TOTAL SSE  (Solver objective -> Min)").font = BOLD
    o = ws.cell(tr, 2, "=" + "+".join(f"{L(c0+1)}{sr}" for c0, _, _ in blocks))
    o.fill, o.border, o.number_format = FILL_O, THIN, "0.00000"

    # van 't Hoff
    vh = RES["Q3"][gas]["vanthoff_toth"]
    vr = tr + 2
    ws.cell(vr, 1, "van 't Hoff:  ln K = ln K0 - dHads/(R*T)").font = BOLD
    for j, h in enumerate(["T (K)", "K (1/MPa)", "1/T (1/K)", "ln K"]):
        c = ws.cell(vr + 1, 1 + j, h)
        c.font, c.fill, c.border = HDR, FILL_H, THIN
    for i, T in enumerate(TS):
        rw = vr + 2 + i
        ws.cell(rw, 1, T)
        ws.cell(rw, 2, f"={refs[f'K{T}  (1/MPa)']}").number_format = "0.0000"
        ws.cell(rw, 3, f"=1/A{rw}").number_format = "0.00000000"
        ws.cell(rw, 4, f"=LN(B{rw})").number_format = "0.00000"
        for j in range(1, 5):
            ws.cell(rw, j).border = THIN
    f_, l_ = vr + 2, vr + 4
    ws.cell(vr + 6, 1, "slope = -dHads/R").font = HDR
    ws.cell(vr + 6, 2, f"=SLOPE(D{f_}:D{l_},C{f_}:C{l_})").number_format = "0.00"
    ws.cell(vr + 7, 1, "intercept = ln K0").font = HDR
    ws.cell(vr + 7, 2, f"=INTERCEPT(D{f_}:D{l_},C{f_}:C{l_})").number_format = "0.0000"
    ws.cell(vr + 8, 1, "R^2 of van 't Hoff line").font = HDR
    ws.cell(vr + 8, 2, f"=RSQ(D{f_}:D{l_},C{f_}:C{l_})").number_format = "0.00000"
    ws.cell(vr + 9, 1, "R (J/mol/K)").font = HDR
    ws.cell(vr + 9, 2, 8.314)
    ws.cell(vr + 10, 1, "dHads  (kJ/mol)").font = BOLD
    c = ws.cell(vr + 10, 2, f"=-B{vr+6}*B{vr+9}/1000")
    c.fill, c.border, c.number_format = FILL_O, THIN, "0.00"
    ws.cell(vr + 10, 3, f"<- exothermic, {vh['dH_kJ']:.1f} kJ/mol").font = NOTE

    for bi, (c0, f_2, l_2) in enumerate(blocks):
        scatter(ws, f"{L(17 + bi*8)}4", f"{gas} @ {TS[bi]} K (Toth)",
                Reference(ws, min_col=c0, min_row=f_2, max_row=l_2),
                Reference(ws, min_col=c0 + 1, min_row=hr, max_row=l_2),
                Reference(ws, min_col=c0 + 2, min_row=hr, max_row=l_2),
                "P (MPa)", "q (mol/kg)")
    return ws


# =====================================================================
# Q3 - multi-site Langmuir
# =====================================================================
def sheet_msl(wb, gas, sets):
    ws = wb.create_sheet(f"Q3_{gas}_MSL")
    widths(ws, {"A": 26, **{get_column_letter(c): 12 for c in range(2, 22)}})
    r = RES["Q3"][gas]["msl"]
    title(ws, f"Q3 - Multi-site Langmuir, {gas} on zeolite 13X", 12)
    note(ws, 2, "Model:  theta = q/qmax = K*p*(1-theta)^a  -  implicit in q, so it is regressed "
                "in the pressure direction:  theta = q_exp/qmax is known, hence "
                "p_calc = theta/(K*(1-theta)^a) is explicit.  Objective = SUM(ln p_calc - ln p_exp)^2.", 12)
    note(ws, 3, "qmax is FIXED at the Toth saturation capacity - qmax and a are strongly "
                "correlated and an unconstrained Solver run drives a to unphysical values.  "
                "Solver variables: a and the three K.  The p = 0 point is excluded (ln 0).", 12)

    names = ["qmax  (mol/kg)  FIXED", "a  (-)"] + [f"K{T}  (1/MPa)" for T in TS]
    vals = [r["qmax"], r["a"]] + [r["K"][str(T)] for T in TS]
    refs, last = param_block(ws, 5, names, vals)
    ws.cell(6, 2).fill = PatternFill("solid", fgColor="E0E0E0")   # fixed, not a variable

    hr = last + 2
    L = get_column_letter
    blocks = []
    for bi, T in enumerate(TS):
        c0 = 1 + bi * 6
        ws.cell(hr - 1, c0, f"T = {T} K").font = BOLD
        ws.cell(hr - 1, c0).fill = FILL_H
        for j, h in enumerate(["p_exp (MPa)", "q_exp", "theta", "p_calc", "ln resid", "sq"]):
            c = ws.cell(hr, c0 + j, h)
            c.font, c.fill, c.border = HDR, FILL_H, THIN
        data = [d for d in sets[T] if d[0] > 0]
        for i, (p, q) in enumerate(data):
            rw = hr + 1 + i
            A, B, C, D, E = (L(c0 + k) for k in range(5))
            ws.cell(rw, c0, p).number_format = "0.00000"
            ws.cell(rw, c0 + 1, q).number_format = "0.000"
            ws.cell(rw, c0 + 2, f"={B}{rw}/{refs['qmax  (mol/kg)  FIXED']}").number_format = "0.0000"
            ws.cell(rw, c0 + 3,
                    f"={C}{rw}/({refs[f'K{T}  (1/MPa)']}*(1-{C}{rw})^{refs['a  (-)']})"
                    ).number_format = "0.00000"
            ws.cell(rw, c0 + 4, f"=LN({D}{rw})-LN({A}{rw})").number_format = "0.0000"
            ws.cell(rw, c0 + 5, f"={E}{rw}^2").number_format = "0.00E+00"
            for j in range(6):
                ws.cell(rw, c0 + j).border = THIN
        blocks.append((c0, hr + 1, hr + len(data)))

    sr = hr + max(len([d for d in sets[T] if d[0] > 0]) for T in TS) + 2
    for bi, (c0, f_, l_) in enumerate(blocks):
        ws.cell(sr, c0, f"SSE(ln p) {TS[bi]} K").font = HDR
        ws.cell(sr, c0 + 1, f"=SUM({L(c0+5)}{f_}:{L(c0+5)}{l_})").number_format = "0.00000"
    tr = sr + 2
    ws.cell(tr, 1, "TOTAL SSE  (Solver objective -> Min)").font = BOLD
    o = ws.cell(tr, 2, "=" + "+".join(f"{L(c0+1)}{sr}" for c0, _, _ in blocks))
    o.fill, o.border, o.number_format = FILL_O, THIN, "0.00000"
    ws.cell(tr + 1, 1, "R^2 in q (from Python check)").font = HDR
    ws.cell(tr + 1, 2, round(r["R2"], 5))

    vh = RES["Q3"][gas]["vanthoff_msl"]
    vr = tr + 3
    ws.cell(vr, 1, "van 't Hoff:  ln K = ln K0 - dHads/(R*T)").font = BOLD
    for j, h in enumerate(["T (K)", "K (1/MPa)", "1/T (1/K)", "ln K"]):
        c = ws.cell(vr + 1, 1 + j, h)
        c.font, c.fill, c.border = HDR, FILL_H, THIN
    for i, T in enumerate(TS):
        rw = vr + 2 + i
        ws.cell(rw, 1, T)
        ws.cell(rw, 2, f"={refs[f'K{T}  (1/MPa)']}").number_format = "0.0000"
        ws.cell(rw, 3, f"=1/A{rw}").number_format = "0.00000000"
        ws.cell(rw, 4, f"=LN(B{rw})").number_format = "0.00000"
        for j in range(1, 5):
            ws.cell(rw, j).border = THIN
    f_, l_ = vr + 2, vr + 4
    ws.cell(vr + 6, 1, "slope = -dHads/R").font = HDR
    ws.cell(vr + 6, 2, f"=SLOPE(D{f_}:D{l_},C{f_}:C{l_})").number_format = "0.00"
    ws.cell(vr + 7, 1, "R^2 of van 't Hoff line").font = HDR
    ws.cell(vr + 7, 2, f"=RSQ(D{f_}:D{l_},C{f_}:C{l_})").number_format = "0.00000"
    ws.cell(vr + 8, 1, "R (J/mol/K)").font = HDR
    ws.cell(vr + 8, 2, 8.314)
    ws.cell(vr + 9, 1, "dHads  (kJ/mol)").font = BOLD
    c = ws.cell(vr + 9, 2, f"=-B{vr+6}*B{vr+8}/1000")
    c.fill, c.border, c.number_format = FILL_O, THIN, "0.00"
    ws.cell(vr + 9, 3, f"<- {vh['dH_kJ']:.1f} kJ/mol").font = NOTE
    return ws


# =====================================================================
def sheet_summary(wb):
    ws = wb.create_sheet("Summary", 0)
    widths(ws, {"A": 34, "B": 16, "C": 16, "D": 16, "E": 16, "F": 12})
    title(ws, "Tutorial 2 (CL24303, MO2026) - Q2 & Q3 regression results", 6)
    note(ws, 2, "All numbers below are the converged Solver values reproduced on the individual "
                "sheets.  Yellow cells = Solver variables, green cell = Solver objective.")

    r = 4
    ws.cell(r, 1, "Q2 - Langmuir,  q = qmax*K*p/(1+K*p),  CO2 at 25 °C").font = BOLD
    for j, h in enumerate(["Adsorbent", "qmax (mmol/g)", "K (1/mmHg)", "R^2", "RMSE", "AARD %"]):
        c = ws.cell(r + 1, 1 + j, h)
        c.font, c.fill, c.border = HDR, FILL_H, THIN
    for i, (lab, key) in enumerate([("Activated carbon", "carbon"), ("Silica gel", "silica")]):
        d = RES["Q2"][key]
        row = [lab, d["qmax"], d["K"], d["R2"], d["RMSE"], d["AARD"]]
        for j, v in enumerate(row):
            c = ws.cell(r + 2 + i, 1 + j, v)
            c.border = THIN
            if j: c.number_format = "0.0000" if j != 2 else "0.000000"

    r = 9
    ws.cell(r, 1, "Q3 - Toth,  q = qmax*K*p/[1+(K*p)^n]^(1/n),  zeolite 13X").font = BOLD
    hdr = ["Gas", "qmax (mol/kg)", "n (-)", "K298", "K308", "K323", "R^2", "dHads (kJ/mol)"]
    for j, h in enumerate(hdr):
        c = ws.cell(r + 1, 1 + j, h)
        c.font, c.fill, c.border = HDR, FILL_H, THIN
    for i, g in enumerate(["CO2", "CH4"]):
        t, v = RES["Q3"][g]["toth"], RES["Q3"][g]["vanthoff_toth"]
        row = [g, t["qmax"], t["n"]] + [t["K"][str(T)] for T in TS] + [t["R2"], v["dH_kJ"]]
        for j, val in enumerate(row):
            c = ws.cell(r + 2 + i, 1 + j, val)
            c.border = THIN
            if j: c.number_format = "0.0000"

    r = 14
    ws.cell(r, 1, "Q3 - Multi-site Langmuir,  theta = K*p*(1-theta)^a,  zeolite 13X").font = BOLD
    hdr = ["Gas", "qmax (mol/kg)", "a (-)", "K298", "K308", "K323", "R^2", "dHads (kJ/mol)"]
    for j, h in enumerate(hdr):
        c = ws.cell(r + 1, 1 + j, h)
        c.font, c.fill, c.border = HDR, FILL_H, THIN
    for i, g in enumerate(["CO2", "CH4"]):
        m, v = RES["Q3"][g]["msl"], RES["Q3"][g]["vanthoff_msl"]
        row = [g, m["qmax"], m["a"]] + [m["K"][str(T)] for T in TS] + [m["R2"], v["dH_kJ"]]
        for j, val in enumerate(row):
            c = ws.cell(r + 2 + i, 1 + j, val)
            c.border = THIN
            if j: c.number_format = "0.0000"

    note(ws, 19, "How to re-run in Excel:  Data -> Solver.  Set Objective = the green SSE cell, "
                 "To: Min, By Changing = the yellow parameter cells, Method = GRG Nonlinear, "
                 "tick 'Make Unconstrained Variables Non-Negative', then Solve.")
    note(ws, 20, "Q2: Langmuir describes silica gel well (R^2 = %.4f) but fits activated carbon "
                 "poorly (R^2 = %.3f) - the carbon isotherm is energetically heterogeneous and "
                 "needs a Toth/Freundlich form."
                 % (RES["Q2"]["silica"]["R2"], RES["Q2"]["carbon"]["R2"]))
    note(ws, 21, "Q3: |dHads| for CO2 (~%.0f kJ/mol) is about %.1f x that of CH4 (~%.0f kJ/mol) - "
                 "the quadrupolar CO2 interacts far more strongly with the Na+ cations of 13X, "
                 "which is why 13X separates CO2 from CH4."
                 % (abs(RES["Q3"]["CO2"]["vanthoff_toth"]["dH_kJ"]),
                    RES["Q3"]["CO2"]["vanthoff_toth"]["dH_kJ"] / RES["Q3"]["CH4"]["vanthoff_toth"]["dH_kJ"],
                    abs(RES["Q3"]["CH4"]["vanthoff_toth"]["dH_kJ"])))
    return ws


if __name__ == "__main__":
    wb = Workbook()
    wb.remove(wb.active)
    sheet_summary(wb)
    sheet_q2(wb, "Q2_ActivatedCarbon", CARBON, "carbon")
    sheet_q2(wb, "Q2_Silica", SILICA, "silica")
    for gas, sets in (("CO2", CO2), ("CH4", CH4)):
        sheet_toth(wb, gas, sets)
        sheet_msl(wb, gas, sets)
    wb.save("Tutorial2_Q2_Q3_Solutions.xlsx")
    print("wrote Tutorial2_Q2_Q3_Solutions.xlsx")
